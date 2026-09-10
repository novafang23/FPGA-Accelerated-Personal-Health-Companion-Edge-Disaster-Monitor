"""
generate_sih_presentation.py
Generates the professional 7-slide presentation for SIH26181:
EdgeGuard / ShrikeFi: AI-Powered Personal Health Companion & Edge Disaster Monitor.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette ────────────────────────────────────────────────────────────
C_NAVY      = RGBColor(15, 23, 42)      # #0F172A (Deep Slate Navy)
C_BLUE_DARK = RGBColor(30, 58, 138)     # #1E3A8A (SIH Royal Blue)
C_BLUE_ACC  = RGBColor(37, 99, 235)     # #2563EB (Vibrant Accent Blue)
C_TEAL      = RGBColor(13, 148, 136)    # #0D9488 (MedTech Teal)
C_RED_WARN  = RGBColor(220, 38, 38)     # #DC2626 (Disaster Alert Red)
C_ORANGE    = RGBColor(234, 88, 12)     # #EA580C (Heatwave Orange)
C_CARD_BG   = RGBColor(248, 250, 252)   # #F8FAFC (Light Card Background)
C_CARD_BORD = RGBColor(226, 232, 240)   # #E2E8F0 (Card Border Gray)
C_TEXT_DARK = RGBColor(30, 41, 59)      # #1E293B (Dark Slate Text)
C_TEXT_MUTED= RGBColor(100, 116, 139)   # #64748B (Muted Text)
C_WHITE     = RGBColor(255, 255, 255)   # White
C_GREEN     = RGBColor(22, 163, 74)     # Success Green

ROOT_DIR = r"C:\Users\abhin\OneDrive\Desktop\verilog"
DOCS_IMAGES = os.path.join(ROOT_DIR, "docs", "images")
OUT_PPTX = os.path.join(ROOT_DIR, "docs", "presentation", "SIH26181_EdgeGuard_Presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]  # Blank slide

def add_header(slide, title_text, category_text="MEDTECH / HARDWARE | PS ID: 26181"):
    """Adds a consistent, professional header to the slide."""
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_NAVY
    top_bar.line.color.rgb = C_NAVY

    tx_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(9.0), Inches(0.3))
    tf_cat = tx_cat.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"SMART INDIA HACKATHON 2026  •  {category_text}"
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_TEAL

    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.40), Inches(9.5), Inches(0.65))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = C_WHITE

    tx_team = slide.shapes.add_textbox(Inches(10.5), Inches(0.25), Inches(2.3), Inches(0.65))
    tf_team = tx_team.text_frame
    p_team = tf_team.paragraphs[0]
    p_team.text = "Team CHIPSTERS"
    p_team.alignment = PP_ALIGN.RIGHT
    p_team.font.size = Pt(14)
    p_team.font.bold = True
    p_team.font.color.rgb = C_WHITE
    p_team2 = tf_team.add_paragraph()
    p_team2.text = "Team ID: 01"
    p_team2.alignment = PP_ALIGN.RIGHT
    p_team2.font.size = Pt(11)
    p_team2.font.color.rgb = C_CARD_BORD

def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORD):
    """Adds a clean rounded card container."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card

def add_stat_box(slide, left, top, width, height, stat_value, stat_label, accent_color=C_BLUE_ACC):
    """Adds an eye-catching stat highlight card."""
    box = add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=accent_color)
    tb = slide.shapes.add_textbox(left, top + Inches(0.08), width, height - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = stat_value
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = accent_color

    p2 = tf.add_paragraph()
    p2.text = stat_label
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_TEXT_DARK
    return box

def add_bullet(tf, title, desc, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = title + ": "
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_DARK
    r = p.add_run()
    r.text = desc
    r.font.bold = False
    r.font.color.rgb = C_TEXT_MUTED
    p.space_after = Pt(8)

# =============================================================================
# SLIDE 1: TITLE PAGE (Brand & Hero Render)
# =============================================================================
s1 = prs.slides.add_slide(blank_layout)

bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY
bg1.line.fill.background()

tx_top = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(0.5))
tf_top = tx_top.text_frame
p = tf_top.paragraphs[0]
p.text = "SMART INDIA HACKATHON 2026  •  MEDTECH / HARDWARE  •  PS ID: 26181"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_TEAL

tx_m = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(7.5), Inches(2.2))
tf_m = tx_m.text_frame
tf_m.word_wrap = True
p = tf_m.paragraphs[0]
p.text = "EdgeGuard / ShrikeFi"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = C_WHITE

p2 = tf_m.add_paragraph()
p2.text = "AI-Powered Personal Health Companion & Edge Disaster Monitor"
p2.font.size = Pt(20)
p2.font.color.rgb = RGBColor(191, 219, 254)

tx_sub = s1.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(7.2), Inches(1.5))
tf_sub = tx_sub.text_frame
tf_sub.word_wrap = True
p = tf_sub.paragraphs[0]
p.text = "A secure, offline medical wearable fusing physiological vitals with environmental disaster telemetry on custom silicon to predict fatal heat stroke, toxic air inhalation, and cold-shock hypothermia before collapse."
p.font.size = Pt(13.5)
p.font.color.rgb = C_CARD_BORD

pill1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.3), Inches(2.2), Inches(0.5))
pill1.fill.solid(); pill1.fill.fore_color.rgb = C_BLUE_ACC; pill1.line.fill.background()
p = pill1.text_frame.paragraphs[0]; p.text = "⚡ 20 ns FPGA Timing"; p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = C_WHITE; p.alignment = PP_ALIGN.CENTER

pill2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.4), Inches(5.3), Inches(2.3), Inches(0.5))
pill2.fill.solid(); pill2.fill.fore_color.rgb = C_TEAL; pill2.line.fill.background()
p = pill2.text_frame.paragraphs[0]; p.text = "🧠 91% INT8 TinyML"; p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = C_WHITE; p.alignment = PP_ALIGN.CENTER

pill3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.9), Inches(5.3), Inches(2.3), Inches(0.5))
pill3.fill.solid(); pill3.fill.fore_color.rgb = C_ORANGE; pill3.line.fill.background()
p = pill3.text_frame.paragraphs[0]; p.text = "🏥 92% MIMIC-III Validated"; p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = C_WHITE; p.alignment = PP_ALIGN.CENTER

tx_t = s1.shapes.add_textbox(Inches(1.0), Inches(6.1), Inches(6.0), Inches(0.8))
tf_t = tx_t.text_frame
p = tf_t.paragraphs[0]
p.text = "Team Name: CHIPSTERS    |    Team ID: 01    |    Category: Hardware"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = RGBColor(148, 163, 184)

hero_img = os.path.join(DOCS_IMAGES, "sih_hero_render.jpg")
if os.path.exists(hero_img):
    s1.shapes.add_picture(hero_img, Inches(8.5), Inches(1.5), Inches(4.2), Inches(4.9))

# =============================================================================
# SLIDE 2: PROBLEM VS SOLUTION
# =============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Problem & Solution — The Deadly 'Smartwatch Blind Spot'")

add_card(s2, Inches(0.8), Inches(1.4), Inches(5.6), Inches(4.5), bg_color=C_WHITE, border_color=C_RED_WARN)
header_red = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(0.6))
header_red.fill.solid(); header_red.fill.fore_color.rgb = C_RED_WARN; header_red.line.fill.background()
p = header_red.text_frame.paragraphs[0]; p.text = "❌ The Problem: Why Current Tech Fails in Disasters"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_WHITE

tx_prob = s2.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(5.2), Inches(3.6))
tf_prob = tx_prob.text_frame
tf_prob.word_wrap = True

add_bullet(tf_prob, "Wearables Are Environmentally Blind", "Smartwatches only track fitness steps or resting pulse. They have zero awareness if ambient air is toxic (AQI 500+) or if 48°C heat is causing heat exhaustion.", first=True)
add_bullet(tf_prob, "Environmental Sensors Ignore the Body", "Weather apps report regional city forecasts, not what an individual worker's cardiovascular system is actually enduring at street level.", first=False)
add_bullet(tf_prob, "Cloud Dependency Trap", "Commercial AI health platforms require mobile networks. During severe storms, floods, or grid failures, cell towers collapse and cloud apps die completely.", first=False)
add_bullet(tf_prob, "Software Jitter Destroys HRV", "Operating system scheduling delays on standard CPUs jitter by 5-20 ms, severely distorting heart rate variability (HRV) metrics.", first=False)

add_card(s2, Inches(6.9), Inches(1.4), Inches(5.6), Inches(4.5), bg_color=C_WHITE, border_color=C_TEAL)
header_green = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(1.4), Inches(5.6), Inches(0.6))
header_green.fill.solid(); header_green.fill.fore_color.rgb = C_TEAL; header_green.line.fill.background()
p = header_green.text_frame.paragraphs[0]; p.text = "✅ The Solution: EdgeGuard Heterogeneous Companion"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_WHITE

tx_sol = s2.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.2), Inches(3.6))
tf_sol = tx_sol.text_frame
tf_sol.word_wrap = True

add_bullet(tf_sol, "Biometric & Environmental Fusion", "Simultaneously acquires dual-wavelength optical PPG pulses along with hyper-local ambient temperature, humidity, and laser PM2.5 particulate matter on the wrist.", first=True)
add_bullet(tf_sol, "Cycle-Accurate FPGA Silicon", "Custom Verilog 8-tap filter and peak-detector FSM runs on 50 MHz hardware with 20 ns resolution, completely bypassing OS scheduling jitter.", first=False)
add_bullet(tf_sol, "Zero-Cloud On-Device TinyML", "Deep neural network fits in only 619 bytes SRAM and evaluates risks in 42 µs, providing 100% offline continuous triage during infrastructure blackouts.", first=False)
add_bullet(tf_sol, "Early Preventive Warnings", "Predicts heat stroke collapse, acute pollution-induced vagal suppression, and cold-water immersion shock 15-20 minutes before medical emergencies.", first=False)

ribbon = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.75))
ribbon.fill.solid(); ribbon.fill.fore_color.rgb = RGBColor(238, 242, 255)
ribbon.line.color.rgb = C_BLUE_ACC
p = ribbon.text_frame.paragraphs[0]
p.text = "💡 Core Innovation: Turning dumb wearable sensors into an autonomous, offline clinical life-saver that detects physiological collapse before symptoms turn fatal."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_BLUE_DARK
p.alignment = PP_ALIGN.CENTER

# =============================================================================
# SLIDE 3: TECHNICAL APPROACH & 4-STAGE PIPELINE
# =============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Technical Approach — 4-Stage Autonomous Dataflow Pipeline")

col_w = Inches(2.75)
col_gap = Inches(0.24)
start_x = Inches(0.8)
card_y = Inches(1.45)
card_h = Inches(5.4)

stages = [
    ("Stage 01", "Sensing Layer", "Physical Data Acquisition", C_BLUE_ACC, [
        ("MAX30102 PPG", "Dual Red & IR optical pulse oximeter for arterial blood flow & SpO2."),
        ("BME280 Climate", "Measures ambient temperature & relative humidity for Steadman Heat Index."),
        ("PMSA003 PM2.5", "Laser scattering sensor measuring hazardous fine particulate pollution.")
    ]),
    ("Stage 02", "FPGA RTL Core", "50 MHz Custom Silicon", C_TEAL, [
        ("Dual 8-Tap Filter", "O(1) running-sum noise filter with bit-shift division (>> 3). 0 DSP, 0 BRAM."),
        ("Systolic Peak FSM", "4-state finite state machine with 250 ms refractory lock-out window."),
        ("20 ns Cycle Timer", "Continuous 50 MHz cycle counter latching heartbeat intervals (IBI) with 0 jitter.")
    ]),
    ("Stage 03", "Clinical TinyML", "ARM Cortex / ESP32-S3", C_NAVY, [
        ("mNEWS2 Engine", "Royal College of Physicians early warning triage protocol score."),
        ("Karlen/Elgendi SQI", "Signal Quality Index rejecting motion noise and false beats."),
        ("6->24->16->3 Model", "INT8 deep micro-network running in 42 µs (91.00% validation accuracy).")
    ]),
    ("Stage 04", "Output & Action", "Edge Alerts & Telemetry", C_ORANGE, [
        ("OLED Display", "128x64 offline display of live HR, SpO2, RR, and hazard threat tiers."),
        ("Haptic Alarm", "Vibration motor & emergency buzzer for immediate worker evacuation."),
        ("WiFi / MQTT Mesh", "Optional emergency beacon broadcasting vitals to disaster rescue teams.")
    ])
]

for i, (stg_num, stg_title, stg_sub, col_color, items) in enumerate(stages):
    cx = start_x + i * (col_w + col_gap)
    card = add_card(s3, cx, card_y, col_w, card_h, bg_color=C_WHITE, border_color=col_color)
    
    top_ban = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, card_y, col_w, Inches(0.85))
    top_ban.fill.solid(); top_ban.fill.fore_color.rgb = col_color; top_ban.line.fill.background()
    p1 = top_ban.text_frame.paragraphs[0]; p1.text = stg_num.upper(); p1.font.size = Pt(10); p1.font.bold = True; p1.font.color.rgb = RGBColor(226, 232, 240); p1.alignment = PP_ALIGN.CENTER
    p2 = top_ban.text_frame.add_paragraph(); p2.text = stg_title; p2.font.size = Pt(13); p2.font.bold = True; p2.font.color.rgb = C_WHITE; p2.alignment = PP_ALIGN.CENTER
    
    tx_s = s3.shapes.add_textbox(cx + Inches(0.1), card_y + Inches(0.9), col_w - Inches(0.2), Inches(0.4))
    p = tx_s.text_frame.paragraphs[0]; p.text = stg_sub; p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = col_color; p.alignment = PP_ALIGN.CENTER
    
    tx_b = s3.shapes.add_textbox(cx + Inches(0.15), card_y + Inches(1.35), col_w - Inches(0.3), card_h - Inches(1.5))
    tf_b = tx_b.text_frame
    tf_b.word_wrap = True
    for j, (it_title, it_desc) in enumerate(items):
        p = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
        p.text = "• " + it_title
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_TEXT_DARK
        
        pr = tf_b.add_paragraph()
        pr.text = it_desc
        pr.font.size = Pt(10)
        pr.font.color.rgb = C_TEXT_MUTED
        pr.space_after = Pt(8)

# =============================================================================
# SLIDE 4: HARDWARE SILICON PROOF & TIMING CLOSURE
# =============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Hardware Acceleration — 20 ns Cycle-Accurate Verilog RTL")

box_analog = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.75))
box_analog.fill.solid(); box_analog.fill.fore_color.rgb = RGBColor(239, 246, 255)
box_analog.line.color.rgb = C_BLUE_ACC
p = box_analog.text_frame.paragraphs[0]
p.text = "💡 Layman Analogy: Why FPGA? Software CPUs have OS jitter — like timing a sprint with a loose, flickering stopwatch. An FPGA is a custom physical circuit: it gives us 20 billionths of a second (20 ns) timing accuracy with zero lag."
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = C_BLUE_DARK
p.alignment = PP_ALIGN.CENTER

add_card(s4, Inches(0.8), Inches(2.25), Inches(5.7), Inches(4.8), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_f = s4.shapes.add_textbox(Inches(1.0), Inches(2.35), Inches(5.3), Inches(0.4))
p = tx_f.text_frame.paragraphs[0]; p.text = "⚡ Renesas ForgeFPGA Silicon Synthesis & Floorplan"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

fp_img = os.path.join(DOCS_IMAGES, "forgefpga_chip_schematic.png")
if os.path.exists(fp_img):
    s4.shapes.add_picture(fp_img, Inches(1.0), Inches(2.8), Inches(5.3), Inches(2.4))

box1 = add_stat_box(s4, Inches(1.0), Inches(5.35), Inches(1.65), Inches(1.5), "17.4%", "LUT5 Logic Used\n(195 / 1120)", C_TEAL)
box2 = add_stat_box(s4, Inches(2.8), Inches(5.35), Inches(1.7), Inches(1.5), "0 DSP / 0 BRAM", "Pure Logic Gates\n(Lowest Cost BOM)", C_BLUE_ACC)
box3 = add_stat_box(s4, Inches(4.65), Inches(5.35), Inches(1.65), Inches(1.5), "+5.603 ns", "WNS Timing Slack\n(STA Met at 69 MHz)", C_GREEN)

add_card(s4, Inches(6.8), Inches(2.25), Inches(5.7), Inches(4.8), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_w = s4.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.3), Inches(0.4))
p = tx_w.text_frame.paragraphs[0]; p.text = "📈 Verified GTKWave Cycle-Accurate Timing Simulation"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

wf_img = os.path.join(DOCS_IMAGES, "waveform_snapshot.png")
if os.path.exists(wf_img):
    s4.shapes.add_picture(wf_img, Inches(7.0), Inches(2.8), Inches(5.3), Inches(2.2))

tx_wb = s4.shapes.add_textbox(Inches(7.0), Inches(5.15), Inches(5.3), Inches(1.75))
tf_wb = tx_wb.text_frame
tf_wb.word_wrap = True

add_bullet(tf_wb, "Noise Filtering Convergence", "Raw optical signals are smoothed by the 8-tap pipeline into filter_red_out[7:0] with 0 multiplier delay.", first=True)
add_bullet(tf_wb, "Systolic Peak Detection", "FSM transitions ARMED -> RISING -> PEAK_FOUND with a 250 ms refractory lock-out window.", first=False)
add_bullet(tf_wb, "Interrupt & IBI Latch", "A single-cycle irq_beat fires, latching the 32-bit heartbeat interval count with 20.0 ns resolution.", first=False)

# =============================================================================
# SLIDE 5: CLINICAL INTELLIGENCE & MIMIC-III BENCHMARK
# =============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Clinical Intelligence — 91% TinyML & MIMIC-III ICU Validation")

add_stat_box(s5, Inches(0.8), Inches(1.35), Inches(3.65), Inches(1.35), "91.00%", "Synthetic Multi-Hazard Accuracy\n(10,000 extreme edge condition test vectors)", C_BLUE_ACC)
add_stat_box(s5, Inches(4.8), Inches(1.35), Inches(3.7), Inches(1.35), "92.05%", "Clinical Concordance on MIMIC-III\n(16,387 gold-standard patient ICU records)", C_TEAL)
add_stat_box(s5, Inches(8.85), Inches(1.35), Inches(3.65), Inches(1.35), "619 Bytes", "Ultra-Low SRAM Memory Footprint\n(Executes in 42 µs with zero cloud)", C_GREEN)

add_card(s5, Inches(0.8), Inches(2.9), Inches(5.7), Inches(4.2), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_nn_h = s5.shapes.add_textbox(Inches(1.0), Inches(3.05), Inches(5.3), Inches(0.4))
p = tx_nn_h.text_frame.paragraphs[0]; p.text = "🧠 On-Device TinyML INT8 Micro-Engine"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

tx_nn = s5.shapes.add_textbox(Inches(1.0), Inches(3.5), Inches(5.3), Inches(3.4))
tf_nn = tx_nn.text_frame
tf_nn.word_wrap = True

add_bullet(tf_nn, "Deep Micro-Architecture", "6 Inputs -> 24 Hidden (ReLU) -> 16 Hidden (ReLU) -> 3 Multi-Hazard Outputs (Logistic Sigmoid).", first=True)
add_bullet(tf_nn, "Quantization-Aware Training (QAT)", "Converted to fixed-point INT8 arithmetic. Zero floating-point math prevents battery drain on edge microcontrollers.", first=False)
add_bullet(tf_nn, "Multi-Disaster Risk Outputs", "Independent continuous threat probabilities:\n  1. Heat Stroke Risk (Cardiovascular drift + thermal index)\n  2. Smog Distress Risk (PM2.5 inhalation + HRV vagal drop)\n  3. Flood Hypothermia Risk (Cold shock + bradycardia)", first=False)

add_card(s5, Inches(6.8), Inches(2.9), Inches(5.7), Inches(4.2), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_cl_h = s5.shapes.add_textbox(Inches(7.0), Inches(3.05), Inches(5.3), Inches(0.4))
p = tx_cl_h.text_frame.paragraphs[0]; p.text = "🏥 Peer-Reviewed Clinical Foundations"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

tx_cl = s5.shapes.add_textbox(Inches(7.0), Inches(3.5), Inches(5.3), Inches(3.4))
tf_cl = tx_cl.text_frame
tf_cl.word_wrap = True

add_bullet(tf_cl, "mNEWS2 Triage Protocol", "Royal College of Physicians clinical early warning system scoring respiration, blood oxygen, and heart rate for hospital-grade triage.", first=True)
add_bullet(tf_cl, "Karlen & Elgendi SQI (Signal Quality)", "Algorithms compute skewness, perfusion index, and waveform entropy to eliminate false alarms from arm motion.", first=False)
add_bullet(tf_cl, "Moran's Physiological Strain Index (PSI)", "Gold-standard cardio-thermal formula quantifying heat exhaustion strain on a scale of 0-10.", first=False)
add_bullet(tf_cl, "AHA Brook Autonomic Stress Model", "American Heart Association scientific formula linking acute fine particulate (PM2.5) exposure to vagal heart rate suppression.", first=False)

# =============================================================================
# SLIDE 6: WORKING PROTOTYPE & LIVE DASHBOARD
# =============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Working System — Live GUI Dashboard & Wearable Hardware")

add_card(s6, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.6), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_d_h = s6.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.3), Inches(0.4))
p = tx_d_h.text_frame.paragraphs[0]; p.text = "🖥️ Standalone 60 FPS Desktop GUI Dashboard"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

tx_d = s6.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.3), Inches(4.8))
tf_d = tx_d.text_frame
tf_d.word_wrap = True

add_bullet(tf_d, "Real-Time Optical Oscilloscope", "Visualizes raw & filtered systolic pulses at 60 frames per second using high-performance Win32 GDI graphics (0 external dependencies).", first=True)
add_bullet(tf_d, "Dual-Mode Hardware Connectivity", "Auto-connects to ESP32-S3 USB COM port at 115,200 baud for live streaming, or switches to 6 simulated disaster profiles for offline demonstrations.", first=False)
add_bullet(tf_d, "Comprehensive Digital Biomarkers", "Displays continuous digital readouts: Heart Rate (BPM), SpO2 (%), Respiratory Rate (RPM), RMSSD HRV (ms), and PM2.5 (ug/m3).", first=False)
add_bullet(tf_d, "Live Clinical & TinyML Meters", "Color-coded dynamic meters for Royal College mNEWS2 clinical triage alongside the 3 on-device AI disaster hazard risk scores.", first=False)
add_bullet(tf_d, "One-Click Instant Execution", "Self-compiles with GCC and launches instantly via launch_dashboard.bat from the repository root.", first=False)

add_card(s6, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.6), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_b_h = s6.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(0.4))
p = tx_b_h.text_frame.paragraphs[0]; p.text = "📦 Physical Wearable CAD & Bill of Materials"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

cad_img = os.path.join(DOCS_IMAGES, "sih_exploded_cad.jpg")
if os.path.exists(cad_img):
    s6.shapes.add_picture(cad_img, Inches(7.0), Inches(1.95), Inches(5.3), Inches(2.3))

tx_bom = s6.shapes.add_textbox(Inches(7.0), Inches(4.35), Inches(5.3), Inches(2.5))
tf_bom = tx_bom.text_frame
tf_bom.word_wrap = True

add_bullet(tf_bom, "Custom Ergonomic Enclosure", "Engineered in Autodesk CAD with optical skin aperture, strap clips, and airflow ventilation for thermal & PM2.5 sensors.", first=True)
add_bullet(tf_bom, "KiCad PCB & Routing", "Complete dual-layer hardware PCB designed in hardware/shrikefi/pcb/ for compact Zero PCB / wearable assembly.", first=False)
add_bullet(tf_bom, "Commercial Mass-Production BOM", "• Renesas ForgeFPGA ($1.80) + ESP32-S3 ($2.20)\n• MAX30102 + BME280 + PM2.5 Sensors ($5.80)\n• OLED + Battery + Enclosure ($4.20)\n=> Total Bulk Unit BOM: <$15 (~₹1,250)", first=False)

# =============================================================================
# SLIDE 7: INDIA IMPACT & QUALCOMM DEPLOYMENT ROADMAP
# =============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "National Impact & Production Roadmap (India 2026)")

add_card(s7, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.5), bg_color=C_WHITE, border_color=C_CARD_BORD)
tx_sc_h = s7.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.3), Inches(0.4))
p = tx_sc_h.text_frame.paragraphs[0]; p.text = "🇮🇳 Tailored for Extreme Indian Climate Emergencies"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_NAVY

tx_sc = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.7))
tf_sc = tx_sc.text_frame
tf_sc.word_wrap = True

add_bullet(tf_sc, "☀️ Severe North India Heatwaves (May-June)", "Protects 380+ million construction laborers, agricultural workers, and delivery personnel from fatal heat exhaustion by issuing warnings 15-20 minutes before collapse.", first=True)
add_bullet(tf_sc, "🌫️ Winter Toxic Smog in Indo-Gangetic Plains", "Warns traffic police, elderly individuals, and asthmatic patients when hyper-local PM2.5 triggers dangerous autonomic cardiac vagal suppression.", first=False)
add_bullet(tf_sc, "🌊 Monsoon Flash Floods (Mumbai / Assam)", "Detects cold-water immersion, hyperventilation, and bradycardia when citizens or rescue teams are stranded during electric grid and mobile network blackouts.", first=False)

add_card(s7, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.5), bg_color=C_WHITE, border_color=C_BLUE_ACC)
tx_rd_h = s7.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(0.4))
p = tx_rd_h.text_frame.paragraphs[0]; p.text = "🚀 Commercial Production: Qualcomm Snapdragon Wear"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_DARK

tx_rd = s7.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.7))
tf_rd = tx_rd.text_frame
tf_rd.word_wrap = True

add_bullet(tf_rd, "Hardware Challenge Silicon Roadmap", "While prototyped on Renesas ForgeFPGA + ESP32-S3 for hackathon accessibility, our architecture is designed for direct migration to Qualcomm Snapdragon Wear W5+ Gen 1.", first=True)
add_bullet(tf_rd, "Qualcomm Hexagon™ DSP Acceleration", "The synthesizable Verilog 8-tap filter algorithms map directly to Hexagon Vector eXtensions (HVX) on the Snapdragon Low-Power Island (<5 mW continuous power).", first=False)
add_bullet(tf_rd, "Qualcomm AI Engine (SNPE / QNN)", "The quantized INT8 neural network compiles natively via Snapdragon Neural Processing Engine (SNPE) for microsecond hardware acceleration.", first=False)
add_bullet(tf_rd, "Enterprise & Defense Readiness", "Targeted for National Disaster Response Force (NDRF), municipal sanitation departments, traffic police, and commercial health insurance deployments.", first=False)

# Save the presentation
prs.save(OUT_PPTX)
print(f"[SUCCESS] Presentation generated cleanly at: {OUT_PPTX}")
