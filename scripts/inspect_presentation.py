import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('docs/presentation/FINAL_FINAL_v6.pptx')
print(f"Slide dimensions: {prs.slide_width} x {prs.slide_height} ({prs.slide_width/914400:.2f} in x {prs.slide_height/914400:.2f} in)")

for idx, slide in enumerate(prs.slides):
    print(f"\n=== Slide {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        text_preview = ""
        if shape.has_text_frame:
            text_preview = " | " + shape.text_frame.text[:40].replace('\n', ' ')
        print(f"  Shape {s_idx}: {shape.name} (type={shape.shape_type}) at ({shape.left/914400:.2f}, {shape.top/914400:.2f}) size ({shape.width/914400:.2f} x {shape.height/914400:.2f}){text_preview}")
